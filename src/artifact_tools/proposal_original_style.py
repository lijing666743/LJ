"""Revise the original proposal deck in place, preserving its visual language."""
from copy import deepcopy
import json
from pathlib import Path
import posixpath
import re
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E

from proposal_revision import apply_rich_runs, set_font_size, shape_by_id

ROOT=Path(__file__).resolve().parents[2]
PROJECT=ROOT/'docs/presentations/proposal_revision_20261015'
MATERIAL=ROOT/'sections/defense_20261015'
ASSETS=ROOT/'plots/proposal_revision_20261015/original_style_assets'
FINAL=PROJECT/'exports/李静_开题答辩_原风格修订版.pptx'
DOCX=MATERIAL/'李静_逐页答辩稿_开题版_8至10分钟.docx'
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}

def content():
    data=json.loads((MATERIAL/'revision_content.json').read_text(encoding='utf-8'))
    for s in data['slides']:s['replacements']={}
    replacements={
    4:{49:'多 UAV 通过 U2U 链路协作，实现任务分流与算力互补。'},
    5:{50:'UAV 移动导致邻居关系和链路条件变化\n随机任务到达使卸载与调度更加复杂',
       51:'状态建模：节点特征、链路历史与队列状态\n方法基础：GAT 聚合邻居信息，MAPPO\n学习多智能体协同策略\n研究重点：局部观测下的任务与资源协作',
       52:'核心目标：\n• 提高按时完成任务数量\n• 兼顾主动通信/计算能耗\n• 缓解紧迫工作量',54:'图强化学习\n方法',59:'状态、动作与奖励\n协同策略学习\n目标与代价权衡'},
    7:{23:'[1] Li X, et al. 信用分配与协同学习（UAV-IRS）. IEEE TWC, 2026, 25: 3186–3200.\n[2] Liu Z, et al. 多智能体学习与 UAV 轨迹能效优化. IEEE TWC, 2025, 24(7): 5917–5930.\n[3] Wu L, Du J, Qu J. 城市多 UAV-MEC 的模型辅助 PPO. arXiv:2603.20238, 2026.',
       27:'模型优化与启发式方法',28:'典型方法：凸近似、贪心与规则调度\n优势：模型与约束清晰，可解释性较强\n适用性取决于模型、规模与在线求解开销',
       32:'独立强化学习方法',33:'典型方法：独立 DQN、IPPO\n特点：各智能体分别学习局部策略\n挑战：其他智能体的策略变化影响学习',
       38:'典型方法：QMIX、MAPPO、MADDPG\n特点：利用集中训练支持协同学习\n挑战：动态图、信息权限与动作约束处理',
       42:'1. 相关方法路线',44:'模型优化、独立学习与协同学习\n各有适用条件，需结合场景比较',
       48:'已有方法为动态决策与协同学习\n提供了可借鉴的技术思路',
       50:'3. 本课题关注问题',52:'动态邻居与信息陈旧\n异构资源与时变队列\n多分支动作的联合约束',
       54:'4. 场景与模型差异',56:'所列文献涉及安全通信、轨迹或 RIS\n与本课题的任务队列模型并不相同',
       60:'关注局部历史信息、队列与资源协同\n通过公平对照检验方法作用'},
    8:{18:'动态拓扑与陈旧信息带来协同困难',20:'• 邻居关系与链路随移动变化\n• 决策仅使用本地观测和历史邻居信息',
       23:'• 随机任务到达，任务具有 deadline 约束\n• 等待、传输与计算共同影响完成结果',
       26:'• UAV 的计算、通信与能量预算不同\n• 源端安排卸载，执行端安排 CPU 服务',
       27:'多分支离散动作与多维约束',29:'• 路由、传输队列与资源动作相互耦合\n• 局部掩码与联合执行器保证动作可行'},
    10:{49:'■ 研究场景：固定高度、外生二维移动的异构多 UAV-MEC 网络。UAV 搭载边缘服务器，研究任务分配与资源编排。',
        50:'■ 任务模型：任务在各 UAV 节点按概率随机生成，先进入未绑定队列；可本地计算或单跳 U2U 卸载，不拆分、不转发。',
        51:'■ 控制机制：源 UAV 决定路由、传输队列和无线资源档位；执行 UAV 决定 CPU 队列与频率。新入队任务下一槽服务。',
        54:'有效速率依据资源单元 SINR 与中断门限计算',55:'时延含等待、传输与计算；SINR 同时计入干扰与噪声。',
        56:'核心建模要点：动态拓扑、随机任务、跨槽队列与双侧控制'},
    11:{48:'3-1-2 优化目标与关键研究内容',50:'物理优化目标（主要目标）',
        53:'π 为联合策略，T 为回合槽数，Nₒₙ 为每槽按时完成数；能耗等作为代价与评价量。',
        57:'通信资源、CPU 容量、剩余能量与半双工约束',58:'局部掩码检查单体合法性，联合执行器处理无线冲突与离散降档。',
        62:'任务路由与传输队列、无线资源档位、CPU 队列与频率'},
    12:{46:'3-2 拟采用研究方法',50:'训练期 Critic 利用集中状态；执行期策略仅使用本地观测与已收到的历史邻居信息。',
        52:'感知输入（局部观测）',60:'陈旧 CSI、历史干扰与邻居消息',64:'数据量、计算量、deadline',
        66:'图注意力协同策略网络',85:'多分支离散\n动作解码',106:'路由、传输队列、资源组、宽度\n功率档位、CPU 队列、CPU 频率',
        108:'策略输出（动作提案）',116:'队列选择与频率档位',119:'传输与资源：',120:'队列 / 资源组 / 宽度',124:'离散功率档位',
        125:'奖励与学习目标\n按时完成奖励，扣除过期、紧迫工作量与主动能耗代价\n奖励用于策略学习，效果依据物理指标评价'},
    13:{46:'3-2 拟采用研究方法',59:'纯本地、随机与启发式方法',75:'读取槽初队列、拓扑与历史观测',79:'按掩码生成多分支动作提案',
        81:'联合仲裁：无线冲突、能量检查与降档',83:'执行槽初队列服务，结算完成与过期',85:'计算奖励，槽末到达后形成下一观测',
        87:'收集轨迹，按批次更新 Actor / Critic',89:'预算完成？',91:'冻结策略，按预定协议开展评估',99:'各 UAV 提出局部动作，联合执行器处理冲突；训练采用 CTDE。'},
    15:{51:'针对外生移动、资源能力差异与随机任务到达，建立包含 UAV 状态、任务生命周期、队列、通信和计算资源的系统模型。',
        55:'将时变协作网络表示为动态图，利用节点特征、合法链路历史及其信息时效描述邻居关系。',56:'Gₜ = (V, Eₜ)',
        59:'将各 UAV 视为智能体，依据局部观测联合决定任务路由、通信调度与计算调度，执行器处理跨节点冲突。',
        60:'CTDE：集中训练，局部观测下的策略决策',62:'团队奖励设计技术：',
        63:'围绕按时完成构造训练奖励，纳入过期、紧迫工作量与主动能耗代价，权重及实际效果需通过实验检验。',
        64:'按时完成奖励 − 过期 / 紧迫工作量 / 能耗代价'},
    16:{52:'构建动态异构多 UAV 协同计算模型',60:'面向随机任务与资源竞争，设计基于局部信息的策略，\n研究任务卸载与资源调度的协同。',
        66:'明确源端与执行端的控制权限，\n在通信、计算和能量预算内协调任务与资源。',
        70:'评价方法有效性',72:'通过对照、消融与重复试验评价完成率、\n主动能耗及稳定性，分析适用边界。'},
    19:{27:'恳请老师批评指正'} }
    for s in data['slides']:
        s['replacements']={str(k):v for k,v in replacements.get(s['n'],{}).items()}
    narration={
    1:'各位老师好，我是李静，导师是杨恒老师。我的开题题目是《动态环境下基于强化学习的无人机协同边缘计算研究》。本课题拟研究多架无人机在移动、资源能力不同、任务随机到达的情况下，如何协调本地计算和相互卸载，使有限资源下更多任务按时完成。下面汇报选题背景、研究方案和后续计划。',
    4:'无人机边缘计算能够为应急感知、偏远区域服务等应用提供灵活的计算支持。这些应用中的计算任务往往有截止期，而单架无人机的算力和能量有限，因此需要考虑多架无人机的协作。右图是研究场景的示意，建筑物作为空间与遮挡背景，任务直接在无人机节点随机到达。图中的数据块表示计算任务，橙色箭头表示一次直接卸载，不包含地面终端到无人机的上传过程。',
    5:'多无人机协同面临两方面变化：一方面，移动不断改变邻居关系和链路条件；另一方面，随机任务到达使队列负载持续变化，各节点的资源能力又不相同。因此，不能只看某一时刻的距离或算力来安排卸载。本课题拟利用图注意力组织邻居与链路信息，通过多智能体强化学习学习协作策略。重点关注按时完成的任务数量，同时评价主动通信与计算能耗，以及紧迫工作量的变化。',
    7:'相关研究大致可以从模型优化、独立学习和协同学习三个角度理解。模型优化和启发式方法便于解释约束，适用性与在线开销取决于具体模型。独立学习分别训练各智能体，其他智能体的策略变化会影响学习。协同强化学习利用集中训练提供协同信息，但仍要明确动态图、信息权限和动作约束。所列文献主要提供方法参考，其安全通信、轨迹或反射面场景与本课题有所不同。本课题重点关注陈旧局部信息、跨时隙队列与异构资源共同作用下的任务调度。',
    8:'这里归纳了四个需要解决的问题。首先，拓扑和信道会变化，策略能够获得的信息又可能已经陈旧。其次，任务具有截止期，等待、传输和计算都会消耗服务时间。再次，各无人机的算力、功率和能量预算不同，源端的卸载安排还要与执行端的计算调度配合。最后，决策包含多个离散分支，局部掩码只能检查一部分合法性，半双工等跨节点冲突还需要联合执行器处理。',
    10:'系统拟采用离散时隙模型，无人机在固定高度外生移动，重点研究任务与资源决策。每架无人机每槽按概率产生零或一个新任务，图下方用有任务A、无新任务、有任务B表示一种可能的到达序列。任务数据量和单位数据计算需求随机采样，截止期由参考服务时间与随机裕量确定。任务先进入未绑定队列，再选择本地处理或直接卸载给一架邻居，不拆分、不转发。蓝色箭头表示到达本机队列，橙色箭头表示单跳卸载。源端控制路由和通信资源，执行端选择计算队列与频率。新入队任务下一槽开始服务，时延包含等待过程。',
    11:'系统的主要目标是在通信、计算、能量和信息约束下，提高单位时间内按时完成的任务数量。完成率和主动能耗用于评价不同方法，过期与回合末截断需要分开统计。这里将物理目标与训练奖励区分开：奖励可以加入过期、紧迫工作量和能耗等代价，引导策略学习，但不意味着奖励最大化就获得了全局最优解。决策变量包括任务路由与传输队列、无线资源档位，以及计算队列和频率。',
    12:'拟采用集中训练、分散策略决策的框架。训练阶段的价值网络利用集中状态，执行阶段每架无人机仅依据自身观测和已经收到的邻居历史信息决策。图注意力对不同邻居的信息进行聚合，时序记忆利用历史观测，再通过离散动作解码器生成任务、通信与计算动作提案。动作包含七个分支，但可以按任务路由、传输资源、功率和计算控制四类理解。提案经过执行器处理冲突后才产生实际服务与能耗，团队奖励用于后续学习。',
    13:'技术路线从环境建模开始，先检查任务状态、队列和物理约束，再构建纯本地、随机和启发式基线，随后开展学习方法及对照实验。每个时隙先读取槽初队列、拓扑和历史观测，生成动作提案并进行联合仲裁，然后执行传输与计算。完成和过期结算之后计算奖励，槽末再产生新任务。训练按批次更新网络，到达预设预算后冻结策略并评估，不把预算结束直接称为收敛，也不提前宣称获得最优策略。',
    14:'下面介绍方案涉及的关键技术和拟达到的课题目标。',
    15:'方案主要依靠四项技术支撑。首先建立动态异构系统模型，使任务到达、队列与资源约束具有一致的时序。其次利用图结构表示动态邻居和合法链路历史，避免把陈旧信息当成实时状态。然后通过集中训练支持多无人机协作，输出任务与资源动作提案。最后设计围绕按时完成的团队奖励，同时纳入过期、紧迫工作量和能耗代价。验证时拟与纯本地、随机及启发式方法比较，并通过消融和重复训练检查各项设计的作用。',
    16:'预期完成四方面工作：建立动态异构多无人机协同计算模型，设计协同卸载决策方法，明确双侧资源控制机制，并形成完整的评价方案。实验将关注任务完成、主动能耗和不同场景下的表现，采用一致的预算与信息权限进行比较。研究目标是明确方法的有效条件和不足，具体改善幅度以及稳定性需要由后续实验给出，开题阶段不预设性能结论。',
    18:'工作计划沿用这五个阶段。首先开展文献调研，梳理问题与研究路线。然后完成系统建模，明确状态、动作和优化目标，完善开题报告。接下来搭建仿真环境，逐步实现任务、队列及资源编排方法。在此基础上开展训练、对比和消融实验，分析任务完成与能耗之间的权衡。最后整理研究结果，完成论文写作、修改和答辩准备。'}
    for s in data['slides']:
        if s['n'] in narration:s['notes']=narration[s['n']]
    data['slides'][14]['title']='关键技术';data['slides'][15]['title']='课题目标'
    data['qa']=[
      {'q':'随机任务在图里如何体现','a':'蓝色单向箭头表示任务到达本机队列，下方时间示意给出有任务A、无新任务、有任务B这一种可能序列。每架 UAV 每槽按 Bernoulli 概率产生零或一个新任务，再采样任务数据量与单位数据计算需求，截止期由参考服务时间与随机裕量确定。图中多个卡片表示抽象任务，不表示同一槽到达一批任务，也不是实验数据。'},
      {'q':'任务来源是不是手机或小车','a':'拟采用的抽象模型直接在 UAV 节点生成计算任务，不单独建立手机、小车或地面接入链路。应急感知等属于应用动机，研究主线从任务到达 UAV 队列开始。建筑物只作为空间和遮挡背景。'},
      {'q':'单跳是不是只能画一个方向的箭头','a':'单跳要求任务从源 UAV 直接传到目的 UAV，不经过第三架 UAV。一次任务卸载用一条单向箭头表示。候选邻居关系可以允许两个方向的通信，但它与某次任务的流向不同，图中用无箭头虚线表示邻居关系。'},
      {'q':'为什么不研究轨迹优化','a':'课题将外生移动作为动态环境条件，集中研究任务分配和资源编排。比较方法使用相同移动和任务随机过程，减少其他因素的干扰。轨迹控制和推进能耗不纳入主模型。'},
      {'q':'研究切入点在哪里','a':'拟围绕局部历史信息、截止期队列和异构资源的耦合开展研究，检验图表示与协同决策的作用。GAT 和 MAPPO 是已有技术，具体贡献需要通过文献定位、可复现对照和消融实验进一步明确。'},
      {'q':'双侧控制分别控制什么','a':'源 UAV 控制未绑定任务的路由，并选择已有传输队列、资源组、宽度与功率档位。执行 UAV 选择 CPU 队列和频率，同时安排本地及收到的任务。源 UAV 不能直接替邻居指定计算动作。'},
      {'q':'如何做公平比较','a':'比较方法采用相同环境、任务和移动随机过程、观测权限、训练预算及模型选择规则。训练随机种子与评估场景分开设置，独立测试场景不参与调参，同时报告不同训练种子的结果和波动。'},
      {'q':'能耗降低是否已经有保证','a':'开题阶段不预设改善幅度。完成更多任务可能增加主动能耗，因此计划同时报告按时完成数、完成率、总主动能耗和每完成任务能耗，分析两者权衡。主动能耗只包括发射和 CPU 动态能耗。'}]
    (MATERIAL/'proposal_v2_content.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    return data

def replace_text(shape,text):
    tx=shape.find('p:txBody',NS);paragraphs=tx.findall('a:p',NS);lines=text.split('\n')
    for i,p in enumerate(paragraphs):
        nodes=p.findall('.//a:t',NS)
        if nodes:
            nodes[0].text=lines[i] if i<len(lines) else ''
            for t in nodes[1:]:t.text=''
    if len(lines)>len(paragraphs):
        paragraphs[-1].find('.//a:t',NS).text='\n'.join(lines[len(paragraphs)-1:])

def add_text(root,name,x,y,w,h,text,size=15,color='104495',bold=False,border=False):
    tree=root.find('p:cSld/p:spTree',NS);sid=max(int(e.get('id')) for e in root.findall('.//p:cNvPr',NS))+1
    sp=E.SubElement(tree,'{%s}sp'%NS['p']);nv=E.SubElement(sp,'{%s}nvSpPr'%NS['p']);E.SubElement(nv,'{%s}cNvPr'%NS['p'],id=str(sid),name=name);E.SubElement(nv,'{%s}cNvSpPr'%NS['p'],txBox='1');E.SubElement(nv,'{%s}nvPr'%NS['p'])
    prop=E.SubElement(sp,'{%s}spPr'%NS['p']);xf=E.SubElement(prop,'{%s}xfrm'%NS['a']);E.SubElement(xf,'{%s}off'%NS['a'],x=str(int(x*9525)),y=str(int(y*9525)));E.SubElement(xf,'{%s}ext'%NS['a'],cx=str(int(w*9525)),cy=str(int(h*9525)))
    geo=E.SubElement(prop,'{%s}prstGeom'%NS['a'],prst='roundRect' if border else 'rect');E.SubElement(geo,'{%s}avLst'%NS['a']);E.SubElement(prop,'{%s}noFill'%NS['a'])
    ln=E.SubElement(prop,'{%s}ln'%NS['a'],w='10000');
    if border:
        sf=E.SubElement(ln,'{%s}solidFill'%NS['a']);E.SubElement(sf,'{%s}srgbClr'%NS['a'],val='83A8EE');E.SubElement(ln,'{%s}prstDash'%NS['a'],val='dash')
    else:E.SubElement(ln,'{%s}noFill'%NS['a'])
    tx=E.SubElement(sp,'{%s}txBody'%NS['p']);E.SubElement(tx,'{%s}bodyPr'%NS['a'],wrap='square',anchor='ctr',lIns='25000',rIns='25000',tIns='0',bIns='0');E.SubElement(tx,'{%s}lstStyle'%NS['a'])
    for line in text.split('\n'):
        p=E.SubElement(tx,'{%s}p'%NS['a']);E.SubElement(p,'{%s}pPr'%NS['a'],algn='ctr')
        r=E.SubElement(p,'{%s}r'%NS['a']);rp=E.SubElement(r,'{%s}rPr'%NS['a'],lang='zh-CN',sz=str(int(size*100)),b='1' if bold else '0')
        sf=E.SubElement(rp,'{%s}solidFill'%NS['a']);E.SubElement(sf,'{%s}srgbClr'%NS['a'],val=color);E.SubElement(rp,'{%s}latin'%NS['a'],typeface='Arial');E.SubElement(rp,'{%s}ea'%NS['a'],typeface='微软雅黑');E.SubElement(r,'{%s}t'%NS['a']).text=line
    return sp

def add_line(root,name,x1,y1,x2,y2,color='F28C28',dashed=False,arrow=True):
    tree=root.find('p:cSld/p:spTree',NS);sid=max(int(e.get('id')) for e in root.findall('.//p:cNvPr',NS))+1
    sp=E.SubElement(tree,'{%s}cxnSp'%NS['p']);nv=E.SubElement(sp,'{%s}nvCxnSpPr'%NS['p']);E.SubElement(nv,'{%s}cNvPr'%NS['p'],id=str(sid),name=name);E.SubElement(nv,'{%s}cNvCxnSpPr'%NS['p']);E.SubElement(nv,'{%s}nvPr'%NS['p'])
    prop=E.SubElement(sp,'{%s}spPr'%NS['p']);attrs={}
    if x2<x1:attrs['flipH']='1'
    if y2<y1:attrs['flipV']='1'
    xf=E.SubElement(prop,'{%s}xfrm'%NS['a'],**attrs);E.SubElement(xf,'{%s}off'%NS['a'],x=str(int(min(x1,x2)*9525)),y=str(int(min(y1,y2)*9525)));E.SubElement(xf,'{%s}ext'%NS['a'],cx=str(int(abs(x2-x1)*9525)),cy=str(int(abs(y2-y1)*9525)))
    geo=E.SubElement(prop,'{%s}prstGeom'%NS['a'],prst='line');E.SubElement(geo,'{%s}avLst'%NS['a'])
    ln=E.SubElement(prop,'{%s}ln'%NS['a'],w='21000' if arrow else '12000');sf=E.SubElement(ln,'{%s}solidFill'%NS['a']);E.SubElement(sf,'{%s}srgbClr'%NS['a'],val=color);E.SubElement(ln,'{%s}prstDash'%NS['a'],val='dash' if dashed else 'solid');E.SubElement(ln,'{%s}headEnd'%NS['a'],type='none');E.SubElement(ln,'{%s}tailEnd'%NS['a'],type='triangle' if arrow else 'none',w='sm',len='sm')

def build_ppt(data):
    with ZipFile(PROJECT/'sources/original.pptx') as z:parts={n:z.read(n) for n in z.namelist()}
    for s in data['slides']:
        i=s['n'];path=f'ppt/slides/slide{i}.xml';rt=E.fromstring(parts[path])
        for sid,text in s['replacements'].items():replace_text(shape_by_id(rt,int(sid)),text)
        if i in [4,5,7]:
            sp=shape_by_id(rt,{4:2,5:2,7:3}[i]);sp.getparent().remove(sp)
        if i in [4,5,8,10]:
            rels=E.fromstring(parts[f'ppt/slides/_rels/slide{i}.xml.rels']);targets={r.get('Id'):r.get('Target') for r in rels}
            pic=rt.find('.//p:pic',NS);rid=pic.find('.//a:blip',NS).get('{%s}embed'%NS['r']);media=posixpath.normpath('ppt/slides/'+targets[rid]);parts[media]=(ASSETS/f'scene{i}.png').read_bytes()
        if i==4:
            add_line(rt,'proposal_single_hop_4',789,286,855,286)
            add_text(rt,'proposal_scene_caption',647,612,535,25,'任务在 UAV 节点随机到达，橙色单箭头表示单跳卸载',12)
        if i==5:add_line(rt,'proposal_single_hop_5',119,553,158,525)
        if i==10:
            apply_rich_runs(shape_by_id(rt,53),[('T',0),('n',-25000),(' = (t',0),('n',-25000),('cmp',30000),(' − t',0),('n',-25000),('arr',30000),(') Δt',0)],21)
            set_font_size(shape_by_id(rt,54),15.5);set_font_size(shape_by_id(rt,55),12)
            add_text(rt,'proposal_random_arrival_heading',638,104,566,32,'随机任务到达各 UAV 的本机队列',17,bold=True)
            add_text(rt,'proposal_sender_role',790,151,198,82,'源 UAV 控制\n任务路由与传输队列\n无线资源与功率档位',12.5,border=True)
            add_text(rt,'proposal_executor_role',1011,151,198,82,'执行 UAV 控制\nCPU 队列选择\nCPU 频率档位',12.5,border=True)
            # Arrival arrows point from floating task cards toward their own UAV.
            for n,(x1,y1,x2,y2) in enumerate([(666,366,666,345),(778,335,778,316),(917,369,917,348),(1048,402,1048,381),(1181,377,1181,357)],1):
                add_line(rt,f'proposal_arrival_{n}',x1,y1,x2,y2,'287EEB')
            add_line(rt,'proposal_single_hop_10',707,306,738,292)
            add_text(rt,'proposal_source_label',614,265,90,25,'源 UAV',11.5)
            add_text(rt,'proposal_target_label',739,237,111,25,'目的 UAV',11.5)
            add_text(rt,'proposal_arrival_sequence_heading',650,454,555,24,'同一 UAV 的一种可能到达序列（仅示意）',12,bold=True)
            for n,(x,label) in enumerate([(653,'时隙 t\n有新任务 A'),(840,'时隙 t+1\n无新任务'),(1027,'时隙 t+2\n有新任务 B')],1):
                add_text(rt,f'proposal_arrival_time_{n}',x,481,171,38,label,11.5,border=True)
            add_text(rt,'proposal_task_attribute_label',650,524,555,24,'任务属性：数据量、计算量、截止时间',11.5)
            add_text(rt,'proposal_scene_legend',622,626,600,28,'蓝箭头：任务随机到达本机队列    橙箭头：单跳 U2U 卸载',11.5)
        if i==11:
            apply_rich_runs(shape_by_id(rt,52),[('max   Λ',0),('succ',-25000),('(π) = E',0),('π',-25000),('[Σ',0),('t',-25000),(' N',0),('on',-25000),('(t)] / T',0)],22)
            set_font_size(shape_by_id(rt,53),13.5)
        if i==13:
            # Separate the inherited subtitle from the technical-route label.
            off=shape_by_id(rt,48).find('p:spPr/a:xfrm/a:off',NS)
            off.set('y',str(int(off.get('y'))-9*9525))
        if i==15:
            apply_rich_runs(shape_by_id(rt,52),[('T',0),('n',-25000),(' = (t',0),('n',-25000),('cmp',30000),(' − t',0),('n',-25000),('arr',30000),(') Δt',0)],20)
            for sid in [60,64]:set_font_size(shape_by_id(rt,sid),17)
        if i==19:
            set_font_size(shape_by_id(rt,27),60)
            for t in shape_by_id(rt,14).findall('.//a:t',NS):
                if t.text:t.text=t.text.replace('副标题：','课题：')
        parts[path]=E.tostring(rt,xml_declaration=True,encoding='UTF-8',standalone=True)
    for path in list(parts):
        if re.fullmatch(r'ppt/notesSlides/notesSlide\d+.xml',path):
            rt=E.fromstring(parts[path])
            for sp in rt.findall('.//p:sp',NS):
                ph=sp.find('p:nvSpPr/p:nvPr/p:ph',NS)
                if ph is not None and ph.get('type')=='body':
                    for t in sp.findall('.//a:t',NS):t.text=''
            parts[path]=E.tostring(rt,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)

def build_doc(data):
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    for name,size in [('Normal',11),('Title',19),('Subtitle',10),('Heading 1',14),('Heading 2',11.5)]:
        st=d.styles[name];st.font.name='微软雅黑';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.paragraph_format.space_after=Pt(5)
        for border in st.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    d.styles['Normal'].paragraph_format.line_spacing=1.2
    d.add_paragraph('开题答辩逐页讲稿','Title');d.add_paragraph(data['title'],'Subtitle');d.add_paragraph('对应原风格修订版 19 页 PPT   答辩人 李静   导师 杨恒','Subtitle')
    d.add_paragraph('按约9分钟编排。正文只讲拟研究的问题、方案、目标和计划，不汇报实验进展。章节页一句带过，第10—12页重点解释任务来源、单跳路径与方法。问题准备不计入汇报时间。')
    md=['# 开题答辩逐页讲稿','',data['title'],'']
    for s in data['slides']:
        d.add_paragraph(f"第 {s['n']} 页  {s['title']}",'Heading 2');d.add_paragraph(s['notes']);md+=[f"## 第{s['n']}页 {s['title']}",'',s['notes'],'']
    d.add_paragraph('答辩问题准备','Heading 1')
    for q in data['qa']:
        d.add_paragraph(q['q'],'Heading 2');d.add_paragraph(q['a']);md+=[f"### {q['q']}",'',q['a'],'']
    footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('李静 开题答辩讲稿  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    d.core_properties.title='李静开题答辩逐页讲稿';d.core_properties.author='李静';d.save(DOCX)
    (MATERIAL/'逐页答辩稿_开题版.md').write_text('\n'.join(md),encoding='utf-8')

def validate(data):
    import hashlib
    with ZipFile(FINAL) as z:
        assert len([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+.xml',n)])==19
        alltext=''.join(''.join(E.fromstring(z.read(f'ppt/slides/slide{i}.xml')).xpath('//a:t/text()',namespaces=NS)) for i in range(1,20))
        for forbidden in ['当前实验','开发试验','训练种子试验','A/B/C','40k','seed42','已有开发','4 架 UAV','最优策略','研究法方法','标红','一区不用']:
            assert forbidden not in alltext,forbidden
        with ZipFile(PROJECT/'sources/original.pptx') as src:
            for n in src.namelist():
                if n.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/')):assert src.read(n)==z.read(n)
        for i in [4,5,10]:
            rt=E.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
            arrows=[e for e in rt.findall('.//p:cxnSp',NS) if e.find('p:nvCxnSpPr/p:cNvPr',NS).get('name','').startswith('proposal_single_hop')]
            assert len(arrows)==1
            for edge in arrows:
                assert edge.find('p:spPr/a:ln/a:headEnd',NS).get('type')=='none'
                assert edge.find('p:spPr/a:ln/a:tailEnd',NS).get('type')=='triangle'
    report={'proposal_only':True,'slides':19,'source_master_layout_theme_identical':True,'single_hop_arrows':'one directed task path per scene','notes':'no speech',
      'spoken_characters':sum(len(s['notes']) for s in data['slides']),
      'outputs':{str(p):{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [FINAL,DOCX]}}
    (ROOT/'metadata/audits/proposal_revision_20261015/original_style_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':
    data=content();build_ppt(data);build_doc(data);validate(data)
